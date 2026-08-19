"""Persistent experimental knowledge base for SPPS Planner V4.0.0.

The V4 layer is additive: it does not replace planner calculations.  It stores
real loading/cleavage observations, preserves raw source text, and separates
parsed records from operator-verified records before supervised training.
"""
from __future__ import annotations

from datetime import datetime, timezone
import csv
import json
import math
from pathlib import Path
import re
import sqlite3
import tempfile
from typing import Any, Iterable, Mapping
from uuid import uuid4
import zipfile


SCHEMA_VERSION = 3
STATUSES = ("parsed", "verified", "incomplete", "excluded")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _id() -> str:
    return uuid4().hex


def default_db_path() -> Path:
    try:
        from spps_planner.user_paths import user_file
        return Path(user_file("experimental_v4.sqlite"))
    except Exception:
        return Path.home() / ".spps_planner_public" / "data" / "experimental_v4.sqlite"


def _connect(path: str | Path | None = None) -> sqlite3.Connection:
    destination = Path(path) if path else default_db_path()
    destination.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(destination)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con



def canonical_product_key(value: Any) -> str:
    """Stable product lookup key while preserving the raw label separately.

    Only workbook metadata suffixes and cosmetic separators are ignored. Product
    numbers/letters remain part of the key, so distinct products are never merged.
    """
    text = str(value or "").strip().lower().replace("–", "-").replace("—", "-")
    text = re.sub(r"\s*:\s*\d+(?:\.\d+)?\s*(?:g\s*/\s*mol)?\.?\s*$", "", text, flags=re.I)
    text = re.sub(r"\s*:\s*$", "", text)
    text = re.sub(r"\((?:\d{6,8}|\d{2}[.]\d{2}[.]\d{2}|\d{1,3})\)\s*$", "", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def sequence_signature(sequence: Any) -> tuple[str, tuple[str, ...], str] | None:
    """Parse a sequence into a case-insensitive identity without losing D-form."""
    text = str(sequence or "").strip()
    if not text:
        return None
    try:
        from spps_planner.parser import parse_sequence
        parsed = parse_sequence(text)
        tokens: list[str] = []
        for token in list(parsed.core_tokens or []) + list(getattr(parsed, "branch_tokens", []) or []):
            raw = str(token).strip()
            if raw.lower().startswith("d") and len(raw) > 1:
                tokens.append("d" + raw[1:].upper())
            else:
                tokens.append(raw.upper())
        if not tokens:
            return None
        nterm = str(getattr(parsed, "nterm", "") or "").strip().upper()
        cterm = str(getattr(parsed, "cterm_text", "") or "").strip().upper()
        return nterm, tuple(tokens), cterm
    except Exception:
        return None


def canonical_sequence_key(sequence: Any) -> str:
    signature = sequence_signature(sequence)
    if signature:
        nterm, tokens, cterm = signature
        return "|".join([nterm, *tokens, cterm])
    text = str(sequence or "").strip()
    return re.sub(r"\s+", "", text).upper() if text else ""


def canonical_resin_key(value: Any) -> str:
    return re.sub(r"\s+", "", normalize_resin(value)).casefold()


def canonical_amino_acid_key(value: Any) -> str:
    return re.sub(r"\s+", "", normalize_amino_acid(value)).casefold()


def _ensure_column(con: sqlite3.Connection, table: str, column: str, ddl: str) -> None:
    columns = {str(row[1]) for row in con.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in columns:
        con.execute(f"ALTER TABLE {table} ADD COLUMN {column} {ddl}")


def _backfill_lookup_keys(con: sqlite3.Connection) -> None:
    """Migrate existing DBs without changing any raw historical value."""
    for row in con.execute("SELECT record_id,resin_type,amino_acid_normalized,resin_key,amino_acid_key FROM loading_records").fetchall():
        resin_key = canonical_resin_key(row["resin_type"])
        aa_key = canonical_amino_acid_key(row["amino_acid_normalized"])
        if row["resin_key"] != resin_key or row["amino_acid_key"] != aa_key:
            con.execute("UPDATE loading_records SET resin_key=?,amino_acid_key=? WHERE record_id=?", (resin_key, aa_key, row["record_id"]))
    for row in con.execute("SELECT record_id,product,sequence,product_key,sequence_key FROM cleavage_records").fetchall():
        product_key = canonical_product_key(row["product"])
        sequence_key = canonical_sequence_key(row["sequence"])
        if row["product_key"] != product_key or row["sequence_key"] != sequence_key:
            con.execute("UPDATE cleavage_records SET product_key=?,sequence_key=? WHERE record_id=?", (product_key, sequence_key, row["record_id"]))
    for row in con.execute("SELECT record_id,product,sequence,product_key,sequence_key FROM synthesis_sequence_records").fetchall():
        product_key = canonical_product_key(row["product"])
        sequence_key = canonical_sequence_key(row["sequence"])
        if row["product_key"] != product_key or row["sequence_key"] != sequence_key:
            con.execute("UPDATE synthesis_sequence_records SET product_key=?,sequence_key=? WHERE record_id=?", (product_key, sequence_key, row["record_id"]))

def initialize(path: str | Path | None = None) -> Path:
    destination = Path(path) if path else default_db_path()
    with _connect(destination) as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS import_sources (
                source_id TEXT PRIMARY KEY,
                imported_at TEXT NOT NULL,
                source_name TEXT NOT NULL,
                source_kind TEXT NOT NULL,
                source_path TEXT NOT NULL,
                sha256 TEXT NOT NULL DEFAULT '',
                note TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS loading_records (
                record_id TEXT PRIMARY KEY,
                status TEXT NOT NULL,
                date TEXT NOT NULL DEFAULT '',
                resin_type TEXT NOT NULL DEFAULT '',
                resin_key TEXT NOT NULL DEFAULT '',
                resin_note TEXT NOT NULL DEFAULT '',
                amino_acid_raw TEXT NOT NULL DEFAULT '',
                amino_acid_normalized TEXT NOT NULL DEFAULT '',
                amino_acid_key TEXT NOT NULL DEFAULT '',
                stereochemistry TEXT NOT NULL DEFAULT '',
                protecting_group TEXT NOT NULL DEFAULT '',
                aa_eq REAL,
                base TEXT NOT NULL DEFAULT '',
                base_eq REAL,
                coupling_reagent TEXT NOT NULL DEFAULT '',
                coupling_reagent_eq REAL,
                additive TEXT NOT NULL DEFAULT '',
                additive_eq REAL,
                loading_time_h REAL,
                loading_solvent TEXT NOT NULL DEFAULT '',
                capping_performed INTEGER,
                capping_method TEXT NOT NULL DEFAULT '',
                resin_sample_weight_mg REAL,
                absorbance REAL,
                loading_rate_mmol_g REAL,
                raw_note TEXT NOT NULL DEFAULT '',
                source_id TEXT,
                source_locator TEXT NOT NULL DEFAULT '',
                outlier_flag INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(source_id) REFERENCES import_sources(source_id)
            );

            CREATE TABLE IF NOT EXISTS cleavage_records (
                record_id TEXT PRIMARY KEY,
                status TEXT NOT NULL,
                product TEXT NOT NULL DEFAULT '',
                product_key TEXT NOT NULL DEFAULT '',
                sequence TEXT NOT NULL DEFAULT '',
                sequence_key TEXT NOT NULL DEFAULT '',
                scale_mmol REAL,
                operator TEXT NOT NULL DEFAULT '',
                tfa_ml REAL,
                tis_ml REAL,
                water_ml REAL,
                other_scavengers_json TEXT NOT NULL DEFAULT '{}',
                cleavage_eq REAL,
                cleavage_time_h REAL,
                temperature_c REAL,
                ether_ml REAL,
                ether_ratio TEXT NOT NULL DEFAULT '',
                filter_ether_ml REAL,
                filter_speed TEXT NOT NULL DEFAULT '',
                crude_g REAL,
                precipitation_good INTEGER,
                separation_problem INTEGER,
                concentration_recommended INTEGER,
                remove_tis_recommended INTEGER,
                overnight_hardening INTEGER,
                raw_observation TEXT NOT NULL DEFAULT '',
                raw_filter_note TEXT NOT NULL DEFAULT '',
                source_id TEXT,
                source_locator TEXT NOT NULL DEFAULT '',
                outlier_flag INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(source_id) REFERENCES import_sources(source_id)
            );

            CREATE TABLE IF NOT EXISTS synthesis_sequence_records (
                record_id TEXT PRIMARY KEY,
                status TEXT NOT NULL,
                product TEXT NOT NULL DEFAULT '',
                product_key TEXT NOT NULL DEFAULT '',
                sequence TEXT NOT NULL DEFAULT '',
                sequence_key TEXT NOT NULL DEFAULT '',
                source_file TEXT NOT NULL DEFAULT '',
                source_page TEXT NOT NULL DEFAULT '',
                source_locator TEXT NOT NULL DEFAULT '',
                row_basis TEXT NOT NULL DEFAULT '',
                raw_note TEXT NOT NULL DEFAULT '',
                source_id TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(source_id) REFERENCES import_sources(source_id)
            );

            CREATE INDEX IF NOT EXISTS idx_loading_lookup
              ON loading_records(resin_type, amino_acid_normalized, status);
            CREATE INDEX IF NOT EXISTS idx_cleavage_lookup
              ON cleavage_records(product, status);
            CREATE INDEX IF NOT EXISTS idx_sequence_lookup
              ON synthesis_sequence_records(product, status);
            """
        )
        _ensure_column(con, "loading_records", "resin_key", "TEXT NOT NULL DEFAULT ''")
        _ensure_column(con, "loading_records", "amino_acid_key", "TEXT NOT NULL DEFAULT ''")
        _ensure_column(con, "cleavage_records", "product_key", "TEXT NOT NULL DEFAULT ''")
        _ensure_column(con, "cleavage_records", "sequence_key", "TEXT NOT NULL DEFAULT ''")
        _ensure_column(con, "synthesis_sequence_records", "product_key", "TEXT NOT NULL DEFAULT ''")
        _ensure_column(con, "synthesis_sequence_records", "sequence_key", "TEXT NOT NULL DEFAULT ''")
        con.execute("CREATE INDEX IF NOT EXISTS idx_loading_key_lookup ON loading_records(resin_key, amino_acid_key, status)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_cleavage_key_lookup ON cleavage_records(product_key, sequence_key, status)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_sequence_key_lookup ON synthesis_sequence_records(product_key, sequence_key, status)")
        _backfill_lookup_keys(con)
    return destination


def _sha256(path: Path) -> str:
    import hashlib
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _register_source(con: sqlite3.Connection, path: Path, kind: str, note: str = "") -> str:
    digest = _sha256(path) if path.is_file() else ""
    existing = con.execute(
        "SELECT source_id FROM import_sources WHERE sha256=? AND source_kind=?",
        (digest, kind),
    ).fetchone() if digest else None
    if existing:
        return str(existing["source_id"])
    source_id = _id()
    con.execute(
        "INSERT INTO import_sources VALUES (?,?,?,?,?,?,?)",
        (source_id, _now(), path.name, kind, str(path), digest, note),
    )
    return source_id


def _float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not text or text in {"-", "—", "–"}:
        return None
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        number = float(match.group(0))
        return number if math.isfinite(number) else None
    except Exception:
        return None


def _volume_ml(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().lower().replace(" ", "")
    if not text or text in {"-", "—", "–"}:
        return None
    # summed forms such as 800ml+300ml are common in the historical workbook
    parts = re.findall(r"(\d+(?:\.\d+)?)(ml|l)", text)
    if not parts:
        return _float(text)
    total = 0.0
    for number, unit in parts:
        total += float(number) * (1000.0 if unit == "l" else 1.0)
    return total


def _scale_mmol(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).lower().replace(" ", "")
    number = _float(text)
    if number is None:
        return None
    # Some legacy rows say ml although the report is clearly using synthesis scale.
    return number


def _eq_time(value: Any) -> tuple[float | None, float | None]:
    text = str(value or "").lower().replace(" ", "")
    eq_match = re.search(r"(\d+(?:\.\d+)?)eq", text)
    h_match = re.search(r"(\d+(?:\.\d+)?)h", text)
    return (
        float(eq_match.group(1)) if eq_match else None,
        float(h_match.group(1)) if h_match else None,
    )


def normalize_resin(value: Any) -> str:
    text = str(value or "").strip()
    low = text.lower()
    if "wang" in low:
        return "Wang resin"
    if "rink" in low and "amide" in low:
        return "Rink Amide resin"
    if "trityl" in low or "ctc" in low:
        return "Trityl/2-CTC resin"
    return text


AA_ALIASES = {
    "A": "Fmoc-Ala-OH", "R": "Fmoc-Arg(Pbf)-OH", "N": "Fmoc-Asn(Trt)-OH",
    "D": "Fmoc-Asp(OtBu)-OH", "C": "Fmoc-Cys(Trt)-OH", "Q": "Fmoc-Gln(Trt)-OH",
    "E": "Fmoc-Glu(OtBu)-OH", "G": "Fmoc-Gly-OH", "H": "Fmoc-His(Trt)-OH",
    "I": "Fmoc-Ile-OH", "L": "Fmoc-Leu-OH", "K": "Fmoc-Lys(Boc)-OH",
    "M": "Fmoc-Met-OH", "F": "Fmoc-Phe-OH", "P": "Fmoc-Pro-OH",
    "S": "Fmoc-Ser(tBu)-OH", "T": "Fmoc-Thr(tBu)-OH", "W": "Fmoc-Trp(Boc)-OH",
    "Y": "Fmoc-Tyr(tBu)-OH", "V": "Fmoc-Val-OH",
}


def normalize_amino_acid(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    compact = re.sub(r"\s+", "", text)
    compact = compact.replace("Fmoc-", "Fmoc-")
    if len(compact) == 1 and compact.upper() in AA_ALIASES:
        return AA_ALIASES[compact.upper()]
    replacements = {
        "fmoc-arg(pbf)-oh": "Fmoc-Arg(Pbf)-OH",
        "fmoc-asn(trt)-oh": "Fmoc-Asn(Trt)-OH",
        "fmoc-asp(otbu)-oh": "Fmoc-Asp(OtBu)-OH",
        "fmoc-cys(trt)-oh": "Fmoc-Cys(Trt)-OH",
        "fmoc-cys(acm)-oh": "Fmoc-Cys(Acm)-OH",
        "fmoc-gln(trt)-oh": "Fmoc-Gln(Trt)-OH",
        "fmoc-glu(otbu)-oh": "Fmoc-Glu(OtBu)-OH",
        "fmoc-his(trt)-oh": "Fmoc-His(Trt)-OH",
        "fmoc-lys(boc)-oh": "Fmoc-Lys(Boc)-OH",
        "fmoc-ser(tbu)-oh": "Fmoc-Ser(tBu)-OH",
        "fmoc-thr(tbu)-oh": "Fmoc-Thr(tBu)-OH",
        "fmoc-trp(boc)-oh": "Fmoc-Trp(Boc)-OH",
        "fmoc-tyr(tbu)-oh": "Fmoc-Tyr(tBu)-OH",
        "fmoc-hyp(tbu)-oh": "Fmoc-Hyp(tBu)-OH",
        "fmoc-cit-oh": "Fmoc-Cit-OH",
        "fmoc-aeea-oh": "Fmoc-AEEA-OH",
    }
    return replacements.get(compact.lower(), text)


def _stereo(name: str) -> str:
    low = name.lower()
    return "D" if "fmoc-d-" in low or low.startswith("d-") else "L/unspecified"


def _protecting_group(name: str) -> str:
    match = re.search(r"\(([^)]+)\)", name)
    return match.group(1) if match else ""


def _keyword_flags(observation: str, filter_note: str = "") -> dict[str, int | None]:
    text = f"{observation}\n{filter_note}".lower()
    good_precip = any(token in text for token in ("석출 잘", "석출은 잘", "침전이 빨", "가루처럼 잘"))
    bad_precip = any(token in text for token in ("석출이 잘 안", "석출x", "침전 안", "석출 안"))
    separation_problem = any(token in text for token in ("분리가 잘 안", "상등액을 버리기 애매", "상층액을 버리기 애매", "상등액 분리 어려"))
    concentrate = any(token in text for token in ("농축 후", "농축해서", "농축후", "농축 추천"))
    remove_tis = "tis는 빼" in text or "tis 빼" in text
    hardening = any(token in text for token in ("overnight", "딱딱", "케이크처럼", "떡지"))
    return {
        "precipitation_good": 0 if bad_precip else (1 if good_precip else None),
        "separation_problem": 1 if separation_problem else 0 if "분리 잘 됨" in text else None,
        "concentration_recommended": 1 if concentrate else 0,
        "remove_tis_recommended": 1 if remove_tis else 0,
        "overnight_hardening": 1 if hardening else 0,
    }


def _crude_g(text: str) -> float | None:
    matches = re.findall(r"(?:cr(?:ude)?\s*[:=]?\s*)(\d+(?:\.\d+)?)\s*g", text, flags=re.I)
    if matches:
        return float(matches[-1])
    return None


def _insert_cleavage(con: sqlite3.Connection, row: Mapping[str, Any]) -> bool:
    existing = con.execute(
        """SELECT record_id FROM cleavage_records
           WHERE source_id=? AND source_locator=?""",
        (row.get("source_id"), row.get("source_locator", "")),
    ).fetchone()
    if existing:
        return False
    keys = [
        "record_id", "status", "product", "product_key", "sequence", "sequence_key", "scale_mmol", "operator",
        "tfa_ml", "tis_ml", "water_ml", "other_scavengers_json", "cleavage_eq",
        "cleavage_time_h", "temperature_c", "ether_ml", "ether_ratio",
        "filter_ether_ml", "filter_speed", "crude_g", "precipitation_good",
        "separation_problem", "concentration_recommended", "remove_tis_recommended",
        "overnight_hardening", "raw_observation", "raw_filter_note", "source_id",
        "source_locator", "outlier_flag", "created_at", "updated_at",
    ]
    values = [row.get(key) for key in keys]
    con.execute(
        f"INSERT INTO cleavage_records ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})",
        values,
    )
    return True


def import_cleavage_report(path: str | Path, db_path: str | Path | None = None) -> dict[str, Any]:
    from openpyxl import load_workbook
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    initialize(db_path)
    workbook = load_workbook(source, data_only=True, read_only=True)
    inserted = 0
    with _connect(db_path) as con:
        source_id = _register_source(con, source, "cleavage_report_xlsx")
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            index = 0
            while index < len(rows):
                row = rows[index]
                product = str(row[0] or "").strip() if row else ""
                # A report block starts with a product row and is followed by Cleavage/TFA.
                if product and index + 1 < len(rows):
                    next_row = rows[index + 1]
                    if str(next_row[0] or "").strip().lower() == "cleavage" and str(next_row[1] or "").strip().upper() == "TFA":
                        scale = _scale_mmol(row[2] if len(row) > 2 else None)
                        operator = str(row[3] or "").strip() if len(row) > 3 else ""
                        tfa_row = next_row
                        tfa_ml = _volume_ml(tfa_row[2] if len(tfa_row) > 2 else None)
                        cleavage_eq, cleavage_time_h = _eq_time(tfa_row[3] if len(tfa_row) > 3 else None)
                        observation = str(tfa_row[4] or "").strip() if len(tfa_row) > 4 else ""
                        components: dict[str, Any] = {}
                        ether_ml = None; ether_ratio = ""; filter_ether_ml = None; filter_speed = ""; filter_note = ""
                        j = index + 2
                        while j < min(index + 8, len(rows)):
                            current = rows[j]
                            section = str(current[0] or "").strip().lower() if current else ""
                            name = str(current[1] or "").strip() if len(current) > 1 else ""
                            value = current[2] if len(current) > 2 else None
                            detail = str(current[3] or "").strip() if len(current) > 3 else ""
                            note = str(current[4] or "").strip() if len(current) > 4 else ""
                            if not any(v not in (None, "") for v in current[:5]):
                                break
                            if section == "filter":
                                filter_ether_ml = _volume_ml(value)
                                filter_speed = detail
                                filter_note = note
                            elif name.upper() == "TIS":
                                components["TIS"] = _volume_ml(value)
                            elif name.upper() in {"H2O", "WATER"}:
                                components["H2O"] = _volume_ml(value)
                            elif name.lower() == "ether":
                                ether_ml = _volume_ml(value)
                                ether_ratio = detail
                            elif name:
                                components[name] = value
                            j += 1
                        flags = _keyword_flags(observation, filter_note)
                        record = {
                            "record_id": _id(), "status": "parsed", "product": product, "product_key": canonical_product_key(product),
                            "sequence": "", "sequence_key": "", "scale_mmol": scale, "operator": operator,
                            "tfa_ml": tfa_ml, "tis_ml": components.pop("TIS", None),
                            "water_ml": components.pop("H2O", None),
                            "other_scavengers_json": json.dumps(components, ensure_ascii=False, default=str),
                            "cleavage_eq": cleavage_eq, "cleavage_time_h": cleavage_time_h,
                            "temperature_c": None, "ether_ml": ether_ml, "ether_ratio": ether_ratio,
                            "filter_ether_ml": filter_ether_ml, "filter_speed": filter_speed,
                            "crude_g": _crude_g(f"{observation}\n{filter_note}"),
                            **flags,
                            "raw_observation": observation, "raw_filter_note": filter_note,
                            "source_id": source_id, "source_locator": f"{sheet.title}!row{index + 1}",
                            "outlier_flag": 0, "created_at": _now(), "updated_at": _now(),
                        }
                        inserted += int(_insert_cleavage(con, record))
                        index = j
                        continue
                index += 1
    return {"kind": "cleavage", "inserted": inserted, "source": str(source)}


def _mapping_value(row: Mapping[str, Any], *names: str) -> Any:
    lookup = {str(key).strip().lower(): value for key, value in row.items()}
    for name in names:
        if name.lower() in lookup:
            return lookup[name.lower()]
    return None


def _insert_loading(con: sqlite3.Connection, row: Mapping[str, Any]) -> bool:
    existing = con.execute(
        "SELECT record_id FROM loading_records WHERE source_id=? AND source_locator=?",
        (row.get("source_id"), row.get("source_locator", "")),
    ).fetchone()
    if existing:
        return False
    keys = [
        "record_id", "status", "date", "resin_type", "resin_key", "resin_note", "amino_acid_raw",
        "amino_acid_normalized", "amino_acid_key", "stereochemistry", "protecting_group", "aa_eq", "base",
        "base_eq", "coupling_reagent", "coupling_reagent_eq", "additive", "additive_eq",
        "loading_time_h", "loading_solvent", "capping_performed", "capping_method",
        "resin_sample_weight_mg", "absorbance", "loading_rate_mmol_g", "raw_note",
        "source_id", "source_locator", "outlier_flag", "created_at", "updated_at",
    ]
    con.execute(
        f"INSERT INTO loading_records ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})",
        [row.get(key) for key in keys],
    )
    return True


def import_loading_csv(path: str | Path, db_path: str | Path | None = None) -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    initialize(db_path)
    inserted = 0
    with source.open("r", encoding="utf-8-sig", newline="") as handle, _connect(db_path) as con:
        reader = csv.DictReader(handle)
        source_id = _register_source(con, source, "loading_csv")
        for number, raw in enumerate(reader, 2):
            aa_raw = str(_mapping_value(raw, "amino_acid", "amino_acid_raw", "aa") or "").strip()
            aa = normalize_amino_acid(_mapping_value(raw, "amino_acid_normalized") or aa_raw)
            resin_raw = _mapping_value(raw, "resin_type", "resin")
            note = str(_mapping_value(raw, "raw_note", "note", "비고") or "")
            capping_raw = _mapping_value(raw, "capping_performed", "capping")
            capping = None
            if capping_raw is not None and str(capping_raw).strip() != "":
                capping = 1 if str(capping_raw).strip().lower() in {"1", "true", "yes", "y", "있음", "약식"} else 0
            record = {
                "record_id": _id(), "status": str(_mapping_value(raw, "status") or "parsed").lower(),
                "date": str(_mapping_value(raw, "date", "날짜") or ""),
                "resin_type": normalize_resin(resin_raw), "resin_key": canonical_resin_key(resin_raw), "resin_note": str(resin_raw or ""),
                "amino_acid_raw": aa_raw, "amino_acid_normalized": aa, "amino_acid_key": canonical_amino_acid_key(aa),
                "stereochemistry": _stereo(aa), "protecting_group": _protecting_group(aa),
                "aa_eq": _float(_mapping_value(raw, "aa_eq", "loading_aa_eq")),
                "base": str(_mapping_value(raw, "base") or "DIEA"),
                "base_eq": _float(_mapping_value(raw, "base_eq", "diea_eq")),
                "coupling_reagent": str(_mapping_value(raw, "coupling_reagent") or ""),
                "coupling_reagent_eq": _float(_mapping_value(raw, "coupling_reagent_eq")),
                "additive": str(_mapping_value(raw, "additive") or ""),
                "additive_eq": _float(_mapping_value(raw, "additive_eq")),
                "loading_time_h": _float(_mapping_value(raw, "loading_time_h", "time_h")),
                "loading_solvent": str(_mapping_value(raw, "loading_solvent", "solvent") or ""),
                "capping_performed": capping,
                "capping_method": str(_mapping_value(raw, "capping_method") or ""),
                "resin_sample_weight_mg": _float(_mapping_value(raw, "resin_sample_weight_mg", "resin_weight_mg")),
                "absorbance": _float(_mapping_value(raw, "absorbance", "abs")),
                "loading_rate_mmol_g": _float(_mapping_value(raw, "loading_rate_mmol_g", "loading_rate")),
                "raw_note": note, "source_id": source_id, "source_locator": f"row{number}",
                "outlier_flag": int(str(_mapping_value(raw, "outlier_flag") or "0").strip() in {"1", "true", "True"}),
                "created_at": _now(), "updated_at": _now(),
            }
            if record["status"] not in STATUSES:
                record["status"] = "parsed"
            inserted += int(_insert_loading(con, record))
    flag_loading_outliers(db_path)
    return {"kind": "loading", "inserted": inserted, "source": str(source)}



_SEQUENCE_NATURAL = set("ARNDCQEGHILKMFPSTWYV")
_SEQUENCE_CORE_ALIASES = {
    "AEEA": "AEEA", "AHX": "Ahx", "CHA": "Cha", "AIB": "Aib", "NLE": "Nle",
    "ORN": "Orn", "CIT": "Cit", "HYP": "Hyp", "DAB": "Dab", "NAL": "Nal",
    "DPR": "Dpr", "BALA": "bAla", "B-ALA": "bAla", "BETAALA": "bAla",
    "GALA": "gAla", "G-ALA": "gAla", "GAMMAALA": "gAla", "GABA": "gAla",
}
_SEQUENCE_MODIFIER_ALIASES = {
    "AC": "Ac", "ACETYL": "Ac", "BOC": "Boc", "FMOC": "Fmoc", "FITC": "FITC",
    "BIOTIN": "Biotin", "FAM": "FAM", "5-FAM": "5-FAM", "6-FAM": "6-FAM",
    "TAMRA": "TAMRA", "CY3": "CY3", "CY5": "CY5", "CY7": "CY7", "PAL": "Pal",
    "MYR": "Myr", "GAL": "Gal", "NIC": "Nic", "CAF": "Caf", "DOTA": "DOTA",
    "NOTA": "NOTA", "DABCYL": "Dabcyl", "BHQ": "BHQ", "NH2": "NH2",
    "CONH2": "CONH2", "AMIDE": "AMIDE", "COOH": "COOH", "CO2H": "CO2H", "OH": "OH",
}
_SEQUENCE_WORD_AA = {
    "ALA":"A", "ARG":"R", "ASN":"N", "ASP":"D", "CYS":"C", "GLN":"Q", "GLU":"E",
    "GLY":"G", "HIS":"H", "ILE":"I", "LEU":"L", "LYS":"K", "MET":"M", "PHE":"F",
    "PRO":"P", "SER":"S", "THR":"T", "TRP":"W", "TYR":"Y", "VAL":"V",
}
_SEQUENCE_TABLE_TERMINATORS = {
    "AAS", "AA", "AMINO ACIDS", "AMINO ACID", "CHEMICALS", "DATE", "NAME", "EQ", "MW", "MMOL",
}


def _std_sequence_tokens(value: Any) -> list[str]:
    """Return conservative sequence tokens from one Check-table cell.

    A page-local STD cell may contain one residue (``E``), a D-form spelling
    (``(D)F``), or a compact pair such as ``Ac-E`` / ``Boc-W``.  Only known
    Planner vocabulary is accepted; arbitrary text is never converted into a
    sequence.
    """
    if not isinstance(value, str):
        return []
    text = value.strip()
    if not text or text.startswith("=") or len(text) > 35:
        return []
    # Cosmetic leading/trailing dashes are common for terminal markers.
    stripped = text.strip().strip("-").strip()
    compact = re.sub(r"\s+", "", stripped).upper()
    if len(compact) == 1 and compact in _SEQUENCE_NATURAL:
        return [compact]
    if re.fullmatch(r"(?i)d[ARNDCQEGHILKMFPSTWYV]", stripped):
        return ["d" + stripped[-1].upper()]
    bracket_d = re.fullmatch(r"(?i)\(D\)[- ]?([ARNDCQEGHILKMFPSTWYV])", stripped)
    if bracket_d:
        return ["d" + bracket_d.group(1).upper()]
    d_named = re.fullmatch(
        r"\(?D\)?[- ]?(ALA|ARG|ASN|ASP|CYS|GLN|GLU|GLY|HIS|ILE|LEU|LYS|MET|PHE|PRO|SER|THR|TRP|TYR|VAL)",
        compact,
    )
    if d_named:
        return ["d" + _SEQUENCE_WORD_AA[d_named.group(1)]]
    if compact in _SEQUENCE_CORE_ALIASES:
        return [_SEQUENCE_CORE_ALIASES[compact]]
    if compact in _SEQUENCE_MODIFIER_ALIASES:
        return [_SEQUENCE_MODIFIER_ALIASES[compact]]
    if re.fullmatch(r"PEG\d+", compact):
        return [compact]
    # Some real STD tables place two known tokens in one cell (e.g. Ac-E or
    # Boc-W). Split only when every sub-token is independently recognized.
    if "-" in stripped:
        parts = [part for part in stripped.split("-") if part.strip()]
        if len(parts) >= 2:
            nested: list[str] = []
            for part in parts:
                parsed = _std_sequence_tokens(part)
                if len(parsed) != 1:
                    return []
                nested.extend(parsed)
            return nested
    return []


def _std_sequence_token(value: Any) -> str | None:
    """Backward-compatible single-token helper."""
    tokens = _std_sequence_tokens(value)
    return tokens[0] if len(tokens) == 1 else None


def _std_sequence_from_values(values: list[Any], row_number: int, start_index: int, end_index: int) -> tuple[str, str] | None:
    from openpyxl.utils import get_column_letter
    tokens: list[str] = []
    coordinates: list[str] = []
    for index in range(start_index, min(end_index + 1, len(values))):
        cell_tokens = _std_sequence_tokens(values[index])
        if cell_tokens:
            tokens.extend(cell_tokens)
            coordinates.append(f"{get_column_letter(index + 1)}{row_number}")
    modifiers = set(_SEQUENCE_MODIFIER_ALIASES.values())
    structural = [token for token in tokens if token not in modifiers]
    if len(structural) < 2:
        return None
    return "-".join(tokens), ",".join(coordinates)


def extract_synthesis_sequence_records(path: str | Path, *, source_file_label: str = "") -> list[dict[str, Any]]:
    """Extract page-local STD sequences from monthly calculation workbooks.

    Contract: a page's Check table is the source of truth for that page. Repeated
    product names on later pages are separate observations, not sequence conflicts.
    Older sheets that put the product name one row below its sequence are supported
    by choosing the adjacent sequence row with the greater number of recognized tokens.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    source = Path(path)
    workbook = load_workbook(source, data_only=False, read_only=True)
    records: list[dict[str, Any]] = []
    skip_pages = {"틀", "틀2", "서식", "未完", "完", "ETC"}
    try:
        for sheet in workbook.worksheets:
            if sheet.title.strip() in skip_pages:
                continue
            max_rows = min(80, int(sheet.max_row or 80))
            max_cols = min(60, int(sheet.max_column or 60))
            matrix = [list(row) for row in sheet.iter_rows(
                min_row=1, max_row=max_rows, min_col=1, max_col=max_cols, values_only=True
            )]
            anchors: list[tuple[int, int]] = []
            for row_index, row in enumerate(matrix):
                for col_index, value in enumerate(row):
                    if isinstance(value, str) and value.strip().lower() == "check table":
                        anchors.append((row_index, col_index))
            for header_row, product_col in anchors:
                header = matrix[header_row]
                result_col = next((
                    col for col in range(product_col + 1, len(header))
                    if isinstance(header[col], str) and header[col].strip().lower() == "result"
                ), None)
                end_col = (result_col - 1) if result_col is not None else min(len(header) - 1, product_col + 28)
                started = False
                blank_streak = 0
                for row_index in range(header_row + 1, min(len(matrix), header_row + 28)):
                    row = matrix[row_index]
                    raw_product = row[product_col] if product_col < len(row) else None
                    product = raw_product.strip() if isinstance(raw_product, str) else ""
                    if product.upper() in _SEQUENCE_TABLE_TERMINATORS:
                        if started:
                            break
                        continue
                    if not product or product.startswith("="):
                        product = ""
                    same = _std_sequence_from_values(row, row_index + 1, product_col + 1, end_col)
                    previous = (
                        _std_sequence_from_values(matrix[row_index - 1], row_index, product_col + 1, end_col)
                        if row_index > header_row + 1 else None
                    )
                    chosen = None
                    row_basis = ""
                    if product:
                        candidates: list[tuple[int, int, tuple[str, str], str]] = []
                        if same:
                            candidates.append((len(same[0].split("-")), 1, same, "same-row"))
                        if previous:
                            candidates.append((len(previous[0].split("-")), 0, previous, "previous-row"))
                        if candidates:
                            candidates.sort(key=lambda value: (value[0], value[1]), reverse=True)
                            _, _, chosen, row_basis = candidates[0]
                    if product and chosen:
                        sequence, sequence_cells = chosen
                        anchor = f"{get_column_letter(product_col + 1)}{header_row + 1}"
                        product_cell = f"{get_column_letter(product_col + 1)}{row_index + 1}"
                        records.append({
                            "status": "parsed",
                            "product": product,
                            "sequence": sequence,
                            "source_file": str(source_file_label or source.name),
                            "source_page": sheet.title,
                            "source_locator": f"Check table {anchor}; product {product_cell}; sequence {sequence_cells}",
                            "row_basis": row_basis,
                            "raw_note": "Page-local STD sequence from Check table; no cross-page canonicalization.",
                        })
                        started = True
                        blank_streak = 0
                    elif started:
                        if not product and not same:
                            blank_streak += 1
                            if blank_streak >= 3:
                                break
                        else:
                            blank_streak = 0
    finally:
        workbook.close()
    return records


def _insert_sequence_record(
    con: sqlite3.Connection,
    row: Mapping[str, Any],
    *,
    replace_same_provenance: bool = False,
) -> bool:
    if replace_same_provenance:
        # Bundled page-local STD seed is regenerated from the same source
        # workbooks when the parser improves.  The sequence-cell list inside the
        # locator can legitimately change when the parser starts recognizing a
        # previously missed token, so the stable slot is workbook + page + product
        # cell (e.g. ``product A3``), not the full locator text.
        locator = str(row.get("source_locator", "") or "")
        slot_match = re.search(r"\bproduct\s+([A-Z]+\d+)\b", locator, flags=re.I)
        slot = slot_match.group(1).upper() if slot_match else ""
        existing = None
        if slot:
            candidates = con.execute(
                """SELECT record_id, source_locator FROM synthesis_sequence_records
                   WHERE source_file=? AND source_page=?""",
                (row.get("source_file", ""), row.get("source_page", "")),
            ).fetchall()
            for candidate in candidates:
                old_match = re.search(r"\bproduct\s+([A-Z]+\d+)\b", str(candidate["source_locator"] or ""), flags=re.I)
                if old_match and old_match.group(1).upper() == slot:
                    existing = candidate
                    break
        else:
            existing = con.execute(
                """SELECT record_id FROM synthesis_sequence_records
                   WHERE source_file=? AND source_page=? AND source_locator=? AND source_id=?""",
                (row.get("source_file", ""), row.get("source_page", ""), locator, row.get("source_id")),
            ).fetchone()
        if existing:
            con.execute(
                """UPDATE synthesis_sequence_records
                   SET status=?, product=?, product_key=?, sequence=?, sequence_key=?, row_basis=?, raw_note=?, source_id=?, updated_at=?
                   WHERE record_id=?""",
                (
                    row.get("status", "parsed"), row.get("product", ""), canonical_product_key(row.get("product", "")), row.get("sequence", ""), canonical_sequence_key(row.get("sequence", "")),
                    row.get("row_basis", ""), row.get("raw_note", ""), row.get("source_id"),
                    row.get("updated_at") or _now(), existing["record_id"],
                ),
            )
            return False
    existing = con.execute(
        """SELECT record_id FROM synthesis_sequence_records
           WHERE source_id=? AND source_file=? AND source_page=? AND source_locator=? AND product=? AND sequence=?""",
        (row.get("source_id"), row.get("source_file", ""), row.get("source_page", ""), row.get("source_locator", ""), row.get("product", ""), row.get("sequence", "")),
    ).fetchone()
    if existing:
        return False
    keys = [
        "record_id", "status", "product", "product_key", "sequence", "sequence_key", "source_file", "source_page", "source_locator",
        "row_basis", "raw_note", "source_id", "created_at", "updated_at",
    ]
    con.execute(
        f"INSERT INTO synthesis_sequence_records ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})",
        [row.get(key) for key in keys],
    )
    return True


def import_synthesis_workbook(path: str | Path, db_path: str | Path | None = None, *, source_file_label: str = "") -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    initialize(db_path)
    parsed = extract_synthesis_sequence_records(source, source_file_label=source_file_label)
    inserted = 0
    with _connect(db_path) as con:
        source_id = _register_source(
            con, source, "synthesis_check_table_xlsx",
            note="Page-local Check table sequences; repeated products remain separate observations.",
        )
        for raw in parsed:
            now = _now()
            row = dict(raw)
            row["product_key"] = canonical_product_key(row.get("product", ""))
            row["sequence_key"] = canonical_sequence_key(row.get("sequence", ""))
            row.update({"record_id": _id(), "source_id": source_id, "created_at": now, "updated_at": now})
            inserted += int(_insert_sequence_record(con, row))
    return {"kind": "sequence_history", "inserted": inserted, "parsed": len(parsed), "source": str(source)}


def import_sequence_history_csv(path: str | Path, db_path: str | Path | None = None) -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    initialize(db_path)
    inserted = 0
    with source.open("r", encoding="utf-8-sig", newline="") as handle, _connect(db_path) as con:
        reader = csv.DictReader(handle)
        source_id = _register_source(con, source, "sequence_history_csv")
        for number, raw in enumerate(reader, 2):
            product = str(_mapping_value(raw, "product", "product_key") or "").strip()
            sequence = str(_mapping_value(raw, "sequence") or "").strip()
            if not product or not sequence:
                continue
            status = str(_mapping_value(raw, "status") or "parsed").lower()
            if status not in STATUSES:
                status = "parsed"
            now = _now()
            row = {
                "record_id": _id(), "status": status, "product": product, "product_key": canonical_product_key(product), "sequence": sequence, "sequence_key": canonical_sequence_key(sequence),
                "source_file": str(_mapping_value(raw, "source_file") or source.name),
                "source_page": str(_mapping_value(raw, "source_page", "source_sheet") or ""),
                "source_locator": str(_mapping_value(raw, "source_locator") or f"row{number}"),
                "row_basis": str(_mapping_value(raw, "row_basis") or ""),
                "raw_note": str(_mapping_value(raw, "raw_note", "source_note") or ""),
                "source_id": source_id, "created_at": now, "updated_at": now,
            }
            inserted += int(_insert_sequence_record(
                con, row, replace_same_provenance=(source.name.lower() == "synthesis_sequence_history_seed.csv")
            ))
    return {"kind": "sequence_history", "inserted": inserted, "source": str(source)}

def import_path(path: str | Path, db_path: str | Path | None = None, *, source_label: str = "") -> list[dict[str, Any]]:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        try:
            with source.open("r", encoding="utf-8-sig", newline="") as handle:
                header = {str(value or "").strip().lower() for value in next(csv.reader(handle), [])}
        except Exception:
            header = set()
        if {"product", "sequence"}.issubset(header) or {"product_key", "sequence"}.issubset(header):
            return [import_sequence_history_csv(source, db_path)]
        return [import_loading_csv(source, db_path)]
    if suffix in {".xlsx", ".xlsm"}:
        # Cleavage Report remains a dedicated stable schema. Monthly calculation
        # workbooks are parsed only through page-local Check tables; no other cells
        # are guessed into experimental records.
        try:
            from openpyxl import load_workbook
            wb = load_workbook(source, data_only=True, read_only=True)
            first = wb[wb.sheetnames[0]]
            first_values = [str(cell.value or "") for row in first.iter_rows(min_row=1, max_row=10) for cell in row]
            looks_like_cleavage = any("Cleavage" in value for value in first_values) and any("TFA" in value for value in first_values)
            wb.close()
        except Exception:
            looks_like_cleavage = False
        if looks_like_cleavage:
            return [import_cleavage_report(source, db_path)]
        parsed = import_synthesis_workbook(source, db_path, source_file_label=source_label)
        if int(parsed.get("parsed", 0)) > 0:
            return [parsed]
        initialize(db_path)
        with _connect(db_path) as con:
            _register_source(con, source, "historical_workbook", note="No supported Cleavage Report or page-local Check table was found; no guessed rows were created.")
        return [{"kind": "registered_workbook", "inserted": 0, "source": str(source)}]
    if suffix == ".zip":
        results: list[dict[str, Any]] = []
        with tempfile.TemporaryDirectory(prefix="spps_v4_import_") as temp:
            with zipfile.ZipFile(source) as archive:
                for info in archive.infolist():
                    name = Path(info.filename)
                    if info.is_dir() or name.name.startswith("~$") or name.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
                        continue
                    safe = Path(temp) / f"{len(results):04d}_{name.name}"
                    safe.write_bytes(archive.read(info))
                    results.extend(import_path(safe, db_path, source_label=name.as_posix()))
        return results
    raise ValueError(f"Unsupported experimental data file: {source.suffix}")



def data_health(db_path: str | Path | None = None) -> dict[str, Any]:
    """Return auditable data-quality metrics without claiming experimental success."""
    initialize(db_path)
    with _connect(db_path) as con:
        def status_counts(table: str) -> dict[str, int]:
            rows = con.execute(f"SELECT status,COUNT(*) AS n FROM {table} GROUP BY status").fetchall()
            return {str(row["status"]): int(row["n"]) for row in rows}

        loading_total = int(con.execute("SELECT COUNT(*) FROM loading_records").fetchone()[0])
        cleavage_total = int(con.execute("SELECT COUNT(*) FROM cleavage_records").fetchone()[0])
        sequence_total = int(con.execute("SELECT COUNT(*) FROM synthesis_sequence_records").fetchone()[0])
        missing_keys = {
            "loading": int(con.execute("SELECT COUNT(*) FROM loading_records WHERE resin_key='' OR amino_acid_key='' ").fetchone()[0]),
            "cleavage": int(con.execute("SELECT COUNT(*) FROM cleavage_records WHERE product<>'' AND product_key='' ").fetchone()[0]),
            "sequence": int(con.execute("SELECT COUNT(*) FROM synthesis_sequence_records WHERE product_key='' OR sequence_key='' ").fetchone()[0]),
        }
        sequence_products = int(con.execute("SELECT COUNT(DISTINCT product_key) FROM synthesis_sequence_records WHERE product_key<>''").fetchone()[0])
        linked_cleavage = int(con.execute("""
            SELECT COUNT(*) FROM cleavage_records c
            WHERE c.product_key<>'' AND EXISTS (
                SELECT 1 FROM synthesis_sequence_records s WHERE s.product_key=c.product_key
            )
        """).fetchone()[0])
        repeated_loading_groups = int(con.execute("""
            SELECT COUNT(*) FROM (
                SELECT resin_key,amino_acid_key FROM loading_records
                WHERE resin_key<>'' AND amino_acid_key<>''
                GROUP BY resin_key,amino_acid_key HAVING COUNT(*)>=2
            )
        """).fetchone()[0])
        loading_rows = [dict(row) for row in con.execute("""
            SELECT resin_key,amino_acid_key,loading_rate_mmol_g FROM loading_records
            WHERE loading_rate_mmol_g IS NOT NULL AND resin_key<>'' AND amino_acid_key<>''
        """).fetchall()]
        cleavage_rows = [dict(row) for row in con.execute("""
            SELECT product_key,cleavage_eq,cleavage_time_h FROM cleavage_records
            WHERE product_key<>'' AND cleavage_eq IS NOT NULL AND cleavage_time_h IS NOT NULL
        """).fetchall()]
        explicit_edt = 0
        for row in con.execute("SELECT other_scavengers_json FROM cleavage_records").fetchall():
            try:
                parsed = json.loads(str(row[0] or "{}"))
            except Exception:
                parsed = {}
            if isinstance(parsed, dict) and any(str(k).strip().casefold()=="edt" and _float(v) not in (None,0) for k,v in parsed.items()):
                explicit_edt += 1

    # Leave-one-out exact-group loading replay. Report MAE, not a fabricated success rate.
    groups: dict[tuple[str,str], list[float]] = {}
    for row in loading_rows:
        groups.setdefault((row["resin_key"], row["amino_acid_key"]), []).append(float(row["loading_rate_mmol_g"]))
    errors: list[float] = []
    for values in groups.values():
        if len(values) < 2:
            continue
        for index, observed in enumerate(values):
            peers = values[:index] + values[index+1:]
            prediction = sorted(peers)[len(peers)//2]
            errors.append(abs(observed-prediction))
    loading_mae = (sum(errors)/len(errors)) if errors else None

    # Historical eq/time replay agreement within a product key. This measures record
    # consistency only; it is deliberately not presented as a biochemical success rate.
    cgroups: dict[str, list[tuple[float,float]]] = {}
    for row in cleavage_rows:
        cgroups.setdefault(str(row["product_key"]), []).append((float(row["cleavage_eq"]),float(row["cleavage_time_h"])))
    agree=0; evaluated=0
    for values in cgroups.values():
        if len(values)<2:
            continue
        counts: dict[tuple[float,float],int]={}
        for value in values: counts[value]=counts.get(value,0)+1
        mode=max(counts,key=counts.get)
        agree += sum(value==mode for value in values)
        evaluated += len(values)
    cleavage_agreement = (100.0*agree/evaluated) if evaluated else None

    return {
        "counts": {"loading": loading_total, "cleavage": cleavage_total, "sequence": sequence_total},
        "status_counts": {
            "loading": status_counts("loading_records"), "cleavage": status_counts("cleavage_records"),
            "sequence": status_counts("synthesis_sequence_records"),
        },
        "missing_canonical_keys": missing_keys,
        "sequence_products": sequence_products,
        "cleavage_records_linked_to_sequence_product": linked_cleavage,
        "repeated_loading_groups": repeated_loading_groups,
        "explicit_edt_records": explicit_edt,
        "loading_leave_one_out_mae_mmol_g": loading_mae,
        "loading_leave_one_out_evaluated": len(errors),
        "cleavage_eq_time_replay_agreement_pct": cleavage_agreement,
        "cleavage_eq_time_replay_evaluated": evaluated,
    }


def preview_path(path: str | Path) -> dict[str, Any]:
    """Parse an import into an isolated temporary DB and return an audit preview.

    The user's real DB is not modified. The exact same import code is exercised, so
    the preview reflects what will actually be stored instead of a second guessed parser.
    """
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    with tempfile.TemporaryDirectory(prefix="spps_v4_preview_") as temp:
        preview_db = Path(temp) / "preview.sqlite"
        results = import_path(source, preview_db)
        loading = list_records("loading", preview_db)
        cleavage = list_records("cleavage", preview_db)
        sequence = list_records("sequence", preview_db)
        warnings: list[str] = []
        registered = sum(1 for row in results if row.get("kind") == "registered_workbook")
        if registered:
            warnings.append(f"{registered} workbook(s) had no supported structured layout; no guessed records will be created.")
        if source.suffix.lower() in {".xlsx", ".xlsm", ".zip"} and not sequence and not cleavage:
            warnings.append("No page-local STD sequence or Cleavage Report record was recognized.")
        samples: list[dict[str, Any]] = []
        for row in sequence[:250]:
            samples.append({
                "kind": "Sequence STD", "status": row.get("status", ""), "product": row.get("product", ""),
                "sequence": row.get("sequence", ""), "source": row.get("source_file", ""),
                "locator": row.get("source_locator", ""),
            })
        for row in cleavage[:100]:
            samples.append({
                "kind": "Cleavage", "status": row.get("status", ""), "product": row.get("product", ""),
                "sequence": row.get("sequence", ""), "source": row.get("source_locator", ""),
                "locator": f"eq={row.get('cleavage_eq','')} / time={row.get('cleavage_time_h','')}",
            })
        for row in loading[:100]:
            samples.append({
                "kind": "Loading", "status": row.get("status", ""), "product": row.get("resin_type", ""),
                "sequence": row.get("amino_acid_normalized", ""), "source": row.get("source_locator", ""),
                "locator": f"loading={row.get('loading_rate_mmol_g','')}",
            })
        return {
            "source": str(source), "results": results,
            "counts": {"loading": len(loading), "cleavage": len(cleavage), "sequence": len(sequence)},
            "registered_workbooks": registered, "warnings": warnings, "samples": samples,
        }


def list_records(kind: str, db_path: str | Path | None = None, *, statuses: Iterable[str] | None = None) -> list[dict[str, Any]]:
    initialize(db_path)
    table = "loading_records" if kind == "loading" else "cleavage_records" if kind == "cleavage" else "synthesis_sequence_records" if kind == "sequence" else None
    if table is None:
        raise ValueError("kind must be loading, cleavage, or sequence")
    clauses = ""
    params: list[Any] = []
    if statuses:
        values = [value for value in statuses if value in STATUSES]
        if values:
            clauses = f" WHERE status IN ({','.join('?' for _ in values)})"
            params.extend(values)
    with _connect(db_path) as con:
        rows = con.execute(f"SELECT * FROM {table}{clauses} ORDER BY created_at, record_id", params).fetchall()
    return [dict(row) for row in rows]


def set_status(kind: str, record_ids: Iterable[str], status: str, db_path: str | Path | None = None) -> int:
    if status not in STATUSES:
        raise ValueError(f"Unsupported status: {status}")
    table = "loading_records" if kind == "loading" else "cleavage_records" if kind == "cleavage" else "synthesis_sequence_records" if kind == "sequence" else None
    if table is None:
        raise ValueError("kind must be loading, cleavage, or sequence")
    ids = [str(value) for value in record_ids if str(value)]
    if not ids:
        return 0
    with _connect(db_path) as con:
        cur = con.execute(
            f"UPDATE {table} SET status=?, updated_at=? WHERE record_id IN ({','.join('?' for _ in ids)})",
            [status, _now(), *ids],
        )
        return int(cur.rowcount)



def update_record(kind: str, record_id: str, changes: Mapping[str, Any], db_path: str | Path | None = None) -> dict[str, Any]:
    table = "loading_records" if kind == "loading" else "cleavage_records" if kind == "cleavage" else None
    if table is None:
        raise ValueError("kind must be loading or cleavage")
    editable = {
        "loading": {"status", "date", "resin_type", "resin_note", "amino_acid_raw", "amino_acid_normalized", "aa_eq", "base", "base_eq", "coupling_reagent", "coupling_reagent_eq", "additive", "additive_eq", "loading_time_h", "loading_solvent", "capping_performed", "capping_method", "resin_sample_weight_mg", "absorbance", "loading_rate_mmol_g", "raw_note", "outlier_flag"},
        "cleavage": {"status", "product", "sequence", "scale_mmol", "operator", "tfa_ml", "tis_ml", "water_ml", "other_scavengers_json", "cleavage_eq", "cleavage_time_h", "temperature_c", "ether_ml", "ether_ratio", "filter_ether_ml", "filter_speed", "crude_g", "precipitation_good", "separation_problem", "concentration_recommended", "remove_tis_recommended", "overnight_hardening", "raw_observation", "raw_filter_note", "outlier_flag"},
    }[kind]
    clean = {str(key): value for key, value in changes.items() if str(key) in editable}
    if "status" in clean and clean["status"] not in STATUSES:
        raise ValueError(f"Unsupported status: {clean['status']}")
    if not clean:
        raise ValueError("No editable fields were supplied.")
    if kind == "loading":
        if "resin_type" in clean:
            clean["resin_type"] = normalize_resin(clean["resin_type"])
            clean["resin_key"] = canonical_resin_key(clean["resin_type"])
        if "amino_acid_raw" in clean or "amino_acid_normalized" in clean:
            raw_aa = clean.get("amino_acid_normalized") or clean.get("amino_acid_raw") or ""
            clean["amino_acid_normalized"] = normalize_amino_acid(raw_aa)
            clean["amino_acid_key"] = canonical_amino_acid_key(clean["amino_acid_normalized"])
    else:
        if "product" in clean:
            clean["product_key"] = canonical_product_key(clean["product"])
        if "sequence" in clean:
            clean["sequence_key"] = canonical_sequence_key(clean["sequence"])
    clean["updated_at"] = _now()
    with _connect(db_path) as con:
        found = con.execute(f"SELECT record_id FROM {table} WHERE record_id=?", (record_id,)).fetchone()
        if not found:
            raise ValueError("Experimental record was not found.")
        assignments = ",".join(f"{key}=?" for key in clean)
        con.execute(f"UPDATE {table} SET {assignments} WHERE record_id=?", [*clean.values(), record_id])
        row = con.execute(f"SELECT * FROM {table} WHERE record_id=?", (record_id,)).fetchone()
    return dict(row)


def add_record(kind: str, values: Mapping[str, Any], db_path: str | Path | None = None, *, status: str = "verified") -> dict[str, Any]:
    """Insert one operator-entered experimental record without fabricating missing values."""
    if status not in STATUSES:
        raise ValueError(f"Unsupported status: {status}")
    initialize(db_path)
    table = "loading_records" if kind == "loading" else "cleavage_records" if kind == "cleavage" else None
    if table is None:
        raise ValueError("kind must be loading or cleavage")
    editable = {
        "loading": {"date", "resin_type", "resin_note", "amino_acid_raw", "amino_acid_normalized", "stereochemistry", "protecting_group", "aa_eq", "base", "base_eq", "coupling_reagent", "coupling_reagent_eq", "additive", "additive_eq", "loading_time_h", "loading_solvent", "capping_performed", "capping_method", "resin_sample_weight_mg", "absorbance", "loading_rate_mmol_g", "raw_note", "outlier_flag"},
        "cleavage": {"product", "sequence", "scale_mmol", "operator", "tfa_ml", "tis_ml", "water_ml", "other_scavengers_json", "cleavage_eq", "cleavage_time_h", "temperature_c", "ether_ml", "ether_ratio", "filter_ether_ml", "filter_speed", "crude_g", "precipitation_good", "separation_problem", "concentration_recommended", "remove_tis_recommended", "overnight_hardening", "raw_observation", "raw_filter_note", "outlier_flag"},
    }[kind]
    now = _now()
    row = {str(k): v for k, v in values.items() if str(k) in editable}
    row.update({"record_id": _id(), "status": status, "created_at": now, "updated_at": now})
    if kind == "loading":
        raw = str(row.get("amino_acid_raw") or row.get("amino_acid_normalized") or "").strip()
        if raw and not row.get("amino_acid_normalized"):
            row["amino_acid_normalized"] = normalize_amino_acid(raw)
        name = str(row.get("amino_acid_normalized") or raw)
        row.setdefault("stereochemistry", _stereo(name))
        row.setdefault("protecting_group", _protecting_group(name))
        if row.get("resin_type"):
            row["resin_type"] = normalize_resin(row["resin_type"])
        row["resin_key"] = canonical_resin_key(row.get("resin_type", ""))
        row["amino_acid_key"] = canonical_amino_acid_key(row.get("amino_acid_normalized") or raw)
    else:
        row["product_key"] = canonical_product_key(row.get("product", ""))
        row["sequence_key"] = canonical_sequence_key(row.get("sequence", ""))
    keys = list(row)
    with _connect(db_path) as con:
        con.execute(f"INSERT INTO {table} ({','.join(keys)}) VALUES ({','.join('?' for _ in keys)})", [row[k] for k in keys])
        saved = con.execute(f"SELECT * FROM {table} WHERE record_id=?", (row["record_id"],)).fetchone()
    return dict(saved)

def flag_loading_outliers(db_path: str | Path | None = None) -> int:
    """Flag statistical/physical review candidates without deleting any value."""
    rows = list_records("loading", db_path)
    changed = 0
    with _connect(db_path) as con:
        for row in rows:
            value = row.get("loading_rate_mmol_g")
            flag = 1 if value is not None and (float(value) < 0 or float(value) > 2.5) else 0
            # Absorbance/loading transcription mistakes often create implausible >2.5 mmol/g values.
            if flag != int(row.get("outlier_flag") or 0):
                con.execute("UPDATE loading_records SET outlier_flag=?, updated_at=? WHERE record_id=?", (flag, _now(), row["record_id"]))
                changed += 1
    return changed


def sources(db_path: str | Path | None = None) -> list[dict[str, Any]]:
    initialize(db_path)
    with _connect(db_path) as con:
        rows = con.execute("SELECT * FROM import_sources ORDER BY imported_at DESC").fetchall()
    return [dict(row) for row in rows]


__all__ = [
    "SCHEMA_VERSION", "STATUSES", "default_db_path", "initialize", "import_cleavage_report",
    "import_loading_csv", "import_sequence_history_csv", "import_synthesis_workbook", "extract_synthesis_sequence_records", "preview_path", "data_health", "canonical_product_key", "canonical_sequence_key", "canonical_resin_key", "canonical_amino_acid_key",
    "import_path", "list_records", "normalize_amino_acid", "normalize_resin",
    "set_status", "update_record", "add_record", "sources", "flag_loading_outliers",
]
