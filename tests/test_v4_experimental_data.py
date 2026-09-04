from __future__ import annotations
from pathlib import Path
import csv


def _write_loading_csv(path: Path, count: int = 36) -> Path:
    fields = ["status","date","resin_type","amino_acid_raw","amino_acid_normalized","aa_eq","base","base_eq","loading_time_h","loading_rate_mmol_g","raw_note"]
    aas = ["Fmoc-Arg(Pbf)-OH","Fmoc-Lys(Boc)-OH","Fmoc-Gly-OH","Fmoc-Ser(tBu)-OH"]
    with path.open("w", encoding="utf-8", newline="") as fh:
        w=csv.DictWriter(fh, fieldnames=fields); w.writeheader()
        for i in range(count):
            aa=aas[i%len(aas)]; eq=[0.5,1.0,2.0][i%3]; beq=[2.0,3.0,4.0][i%3]
            w.writerow(dict(status="parsed",date=f"2030-01-{(i%28)+1:02d}",resin_type="Trityl/2-CTC resin",amino_acid_raw=aa,amino_acid_normalized=aa,aa_eq=eq,base="DIEA",base_eq=beq,loading_time_h=4,loading_rate_mmol_g=round(0.2+0.08*eq+0.003*i,4),raw_note="synthetic public test fixture"))
    return path


def _write_cleavage_xlsx(path: Path, count: int = 12) -> Path:
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="Cleavage Report"
    row=1
    for i in range(count):
        product=f"Demo Peptide {i+1}"
        scale=100
        eq=30 if i != 0 else 80
        total=scale*eq
        tfa=0.95*total; water=0.05*total
        obs="precipitation good" if i else "석출은 잘 되었으나 상등액과 분리가 잘 안됨"
        ws.cell(row,1,product); ws.cell(row,3,f"{scale}mmol"); ws.cell(row,4,"TEST")
        ws.cell(row+1,1,"Cleavage"); ws.cell(row+1,2,"TFA"); ws.cell(row+1,3,f"{tfa}ml"); ws.cell(row+1,4,f"{eq}eq/3h"); ws.cell(row+1,5,obs)
        ws.cell(row+2,2,"TIS"); ws.cell(row+2,3,"-")
        ws.cell(row+3,2,"H2O"); ws.cell(row+3,3,f"{water}ml")
        if i == 1:
            ws.cell(row+4,2,"Thioanisole"); ws.cell(row+4,3,"10ml")
            ws.cell(row+5,2,"Ether"); ws.cell(row+5,3,"10L"); ws.cell(row+5,4,"1:3")
            ws.cell(row+6,1,"Filter"); ws.cell(row+6,2,"Ether"); ws.cell(row+6,3,"2L"); ws.cell(row+6,4,"normal")
        else:
            ws.cell(row+4,2,"Ether"); ws.cell(row+4,3,"10L"); ws.cell(row+4,4,"1:3")
            ws.cell(row+5,1,"Filter"); ws.cell(row+5,2,"Ether"); ws.cell(row+5,3,"2L"); ws.cell(row+5,4,"normal")
        row += 8
    wb.save(path); return path


def test_loading_csv_import_and_review_status(tmp_path):
    from suite_gui import experimental_data
    db=tmp_path/"exp.sqlite"; source=_write_loading_csv(tmp_path/"loading.csv")
    result=experimental_data.import_loading_csv(source,db); assert result["inserted"]>=30
    rows=experimental_data.list_records("loading",db); assert rows and all(r["status"]=="parsed" for r in rows)
    target=rows[0]["record_id"]; assert experimental_data.set_status("loading",[target],"verified",db)==1


def test_cleavage_report_parser_preserves_observation(tmp_path):
    from suite_gui import experimental_data
    db=tmp_path/"exp.sqlite"; source=_write_cleavage_xlsx(tmp_path/"cleavage.xlsx")
    result=experimental_data.import_cleavage_report(source,db); assert result["inserted"]>=10
    row=next(r for r in experimental_data.list_records("cleavage",db) if r["product"]=="Demo Peptide 1")
    assert row["cleavage_eq"]==80.0 and row["cleavage_time_h"]==3.0 and row["separation_problem"]==1


def test_loading_advisor_uses_similarity_evidence(tmp_path):
    from suite_gui import experimental_data, ml_advisor_v4
    db=tmp_path/"exp.sqlite"; experimental_data.import_loading_csv(_write_loading_csv(tmp_path/"loading.csv"),db)
    result=ml_advisor_v4.loading_advice("Trityl/2-CTC resin","Fmoc-Arg(Pbf)-OH",aa_eq=0.5,base_eq=2,db_path=db,include_parsed=True)
    assert result["prediction"] is not None and result["evidence_count"]>=2


def test_supervised_loading_model_requires_verified_records(tmp_path):
    from suite_gui import experimental_data, ml_advisor_v4
    db=tmp_path/"exp.sqlite"; experimental_data.import_loading_csv(_write_loading_csv(tmp_path/"loading.csv"),db)
    rows=experimental_data.list_records("loading",db); experimental_data.set_status("loading",[r["record_id"] for r in rows[:20]],"verified",db)
    result=ml_advisor_v4.loading_advice("Trityl/2-CTC resin","Fmoc-Lys(Boc)-OH",aa_eq=1,base_eq=3,db_path=db,include_parsed=True)
    assert result["prediction"] is not None


def test_unknown_historical_workbook_is_registered_not_fabricated(tmp_path):
    from openpyxl import Workbook
    from suite_gui import experimental_data
    wb=Workbook(); wb.active["A1"]="free-form operator workbook"; source=tmp_path/"history.xlsx"; wb.save(source)
    db=tmp_path/"exp.sqlite"; result=experimental_data.import_path(source,db)
    assert result[0]["kind"]=="registered_workbook" and experimental_data.list_records("loading",db)==[] and experimental_data.list_records("cleavage",db)==[]


def test_record_edit_preserves_source(tmp_path):
    from suite_gui import experimental_data
    db=tmp_path/"exp.sqlite"; experimental_data.import_loading_csv(_write_loading_csv(tmp_path/"loading.csv"),db)
    row=experimental_data.list_records("loading",db)[0]; source_id=row["source_id"]
    edited=experimental_data.update_record("loading",row["record_id"],{"loading_time_h":5.0,"status":"verified"},db)
    assert edited["loading_time_h"]==5.0 and edited["source_id"]==source_id
