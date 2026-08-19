import csv
from pathlib import Path

from openpyxl import Workbook

from suite_gui import experimental_data, ml_advisor_v4


def _write_sequence_csv(path: Path, product: str, sequence: str) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["status", "product", "sequence", "source_file", "source_page", "source_locator", "row_basis", "raw_note"])
        writer.writeheader()
        writer.writerow({
            "status": "parsed", "product": product, "sequence": sequence,
            "source_file": "synthetic.xlsx", "source_page": "run", "source_locator": "Check table A2",
            "row_basis": "same-row", "raw_note": "test",
        })


def test_check_table_import_preserves_page_local_std_and_legacy_previous_row(tmp_path):
    book = tmp_path / "monthly.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "same-row"
    ws["A2"] = "Check table"; ws["B2"] = 4; ws["C2"] = 3; ws["D2"] = 2; ws["E2"] = "LOAD"
    ws["A3"] = "Product-A: 1234.5"; ws["B3"] = "Ac"; ws["C3"] = "C"; ws["D3"] = "G"; ws["E3"] = "K"
    ws["A4"] = "AAs"; ws["B4"] = "NAME"  # section terminator, not a product
    ws2 = wb.create_sheet("legacy")
    ws2["C1"] = "Check table"; ws2["N1"] = "LOAD"
    ws2["D12"] = "W"; ws2["E12"] = "L"; ws2["F12"] = "G"; ws2["N12"] = "AEEA"
    ws2["C13"] = "Product-X"
    wb.save(book)

    parsed = experimental_data.extract_synthesis_sequence_records(book)
    assert [(row["product"], row["sequence"], row["row_basis"]) for row in parsed] == [
        ("Product-A: 1234.5", "Ac-C-G-K", "same-row"),
        ("Product-X", "W-L-G-AEEA", "previous-row"),
    ]

    db = tmp_path / "exp.sqlite"
    result = experimental_data.import_path(book, db)
    assert result[0]["kind"] == "sequence_history"
    records = experimental_data.list_records("sequence", db, statuses=["parsed"])
    assert len(records) == 2


def test_product_sequence_history_is_many_observations_not_forced_one_to_one(tmp_path):
    db = tmp_path / "exp.sqlite"
    first = tmp_path / "seq1.csv"; second = tmp_path / "seq2.csv"
    _write_sequence_csv(first, "Product-X", "Boc-W-L-G-Q-G-G-G-S-K-Dab-FITC-AEEA")
    _write_sequence_csv(second, "Product-X", "W-L-G-Q-G-G-G-S-K-Dab-AEEA")
    experimental_data.import_sequence_history_csv(first, db)
    experimental_data.import_sequence_history_csv(second, db)
    observations = ml_advisor_v4._product_sequence_observations(db)
    assert len(observations[ml_advisor_v4._normalize_product_key("Product-X")]) == 2


def test_page_std_without_cterm_can_support_same_current_core_with_explicit_cterm(tmp_path):
    db = tmp_path / "exp.sqlite"
    mapping = tmp_path / "seq.csv"
    _write_sequence_csv(mapping, "Product-Cterm: 509.6120", "Ac-K-G-H-K")
    experimental_data.import_sequence_history_csv(mapping, db)
    ok, basis = ml_advisor_v4._product_sequence_supported(
        "Product-Cterm", "Ac-KGHK-NH2", __import__("pandas").DataFrame(), db_path=db
    )
    assert ok is True
    assert "Check table" in basis


def test_real_tfa_water_history_is_not_rejected_or_injected_with_edt(tmp_path):
    db = tmp_path / "exp.sqlite"
    mapping = tmp_path / "seq.csv"
    _write_sequence_csv(mapping, "CysPep", "Ac-C-G-NH2")
    experimental_data.import_sequence_history_csv(mapping, db)
    experimental_data.add_record("cleavage", {
        "product": "CysPep", "sequence": "", "scale_mmol": 1.0,
        "tfa_ml": 9.5, "tis_ml": 0.0, "water_ml": 0.5,
        "other_scavengers_json": "{}", "cleavage_eq": 10.0, "cleavage_time_h": 2.0,
        "raw_observation": "recorded TFA/water condition",
    }, db, status="verified")

    result = ml_advisor_v4.cleavage_recommendation(
        product="CysPep", sequence="Ac-C-G-NH2", resin="Rink Amide", scale_mmol=1.0,
        db_path=db, include_parsed=True,
    )
    rec = result["recommended_condition"]
    assert result["method"] == "sequence-matched observed-condition recommendation"
    assert rec["apply_allowed"] is True
    assert rec["composition_pct"] == {"TFA": 95.0, "Water": 5.0}
    assert "EDT" not in rec["composition_pct"]


def test_chemistry_only_cys_fallback_is_reference_not_ml_apply(tmp_path):
    db = tmp_path / "empty.sqlite"
    experimental_data.initialize(db)
    result = ml_advisor_v4.cleavage_recommendation(
        product="Unknown", sequence="Ac-C-G-NH2", resin="Rink Amide", scale_mmol=1.0,
        db_path=db,
    )
    rec = result["recommended_condition"]
    assert rec is not None
    assert rec["recommendation_kind"] == "CHEMISTRY RULE"
    assert rec["condition_source"] == "chemistry_rule_reference"
    assert rec["apply_allowed"] is False



def test_check_table_parser_recognizes_parenthesized_d_forms_and_combined_std_cells(tmp_path):
    book = tmp_path / "monthly_d.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "std"
    ws["A2"] = "Check table"
    ws["A3"] = "Demo-D: 471.51 g/mol"
    ws["D3"] = "Ac"
    ws["E3"] = "(D)A"
    ws["F3"] = "(D)V"
    ws["G3"] = "(D)L"
    ws["H3"] = " -NH2"
    ws["A6"] = "Amino acids"
    ws2 = wb.create_sheet("combined")
    ws2["A2"] = "Check table"
    ws2["A3"] = "Demo-Peptide-3(8):"
    ws2["D3"] = "Ac-G"
    ws2["E3"] = "H"
    ws2["F3"] = "K"
    ws2["G3"] = " -NH2"
    ws2["A6"] = "Amino acids"
    wb.save(book)

    rows = experimental_data.extract_synthesis_sequence_records(book)
    by_product = {row["product"]: row["sequence"] for row in rows}
    assert by_product["Demo-D: 471.51 g/mol"] == "Ac-dA-dV-dL-NH2"
    assert by_product["Demo-Peptide-3(8):"] == "Ac-G-H-K-NH2"


def test_product_lookup_key_ignores_workbook_metadata_but_not_peptide_number():
    assert ml_advisor_v4._normalize_product_key("Demo-Peptide-3(260720)") == ml_advisor_v4._normalize_product_key("Demo Peptide 3(8): 889.0 g/mol")
    assert ml_advisor_v4._normalize_product_key("DEMO251101") == ml_advisor_v4._normalize_product_key("DEMO_251101: 2427.83")
    assert ml_advisor_v4._normalize_product_key("Demo-Peptide-3") != ml_advisor_v4._normalize_product_key("Demo-Peptide-1")


def test_sequence_matching_is_case_insensitive_and_preserves_d_form_semantics():
    assert ml_advisor_v4._sequence_observation_matches("Ac-G-H-K-NH2", "aC-gHk-nh2")
    assert ml_advisor_v4._sequence_observation_matches("Ac-dA-dV-dL-NH2", "ac-da-dv-dl-nh2")
    assert not ml_advisor_v4._sequence_observation_matches("Ac-A-V-L-NH2", "ac-da-dv-dl-nh2")


def test_sequence_matched_but_incomplete_cleavage_history_is_visible_not_replaced_by_rule(tmp_path):
    db = tmp_path / "exp.sqlite"
    mapping = tmp_path / "seq.csv"
    _write_sequence_csv(mapping, "Demo-Peptide-7(8): 708.9460", "G-Q-P-R")
    experimental_data.import_sequence_history_csv(mapping, db)
    experimental_data.add_record("cleavage", {
        "product": "Demo-Peptide-7(260326)", "sequence": "", "scale_mmol": 400.0,
        "tfa_ml": 5700.0, "tis_ml": None, "water_ml": 300.0,
        "other_scavengers_json": "{}", "cleavage_eq": 30.0, "cleavage_time_h": 3.0,
        "raw_observation": "synthetic incomplete historical record",
    }, db, status="verified")
    result = ml_advisor_v4.cleavage_recommendation(
        product="Demo-Peptide-7", sequence="gqpr", resin="Rink Amide",
        scale_mmol=400.0, db_path=db, include_parsed=True,
    )
    assert result["method"] == "sequence-matched history recognized but incomplete"
    assert result["matched_history_count"] == 1
    assert result["recommended_condition"] is None
    assert len(result["evidence"]) == 1
