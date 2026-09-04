from __future__ import annotations

import csv
from pathlib import Path


def _write_usage_csv(path: Path, rows):
    fields = [
        "status","product_raw","product","product_key","sequence_raw","sequence","sequence_key","scale_mmol","operator","manufacture_period","synthesis_key","record_scope",
        "tfa_raw","tfa_value","tfa_unit","tfa_ml","tfa_status","water_raw","water_value","water_unit","water_ml","water_status","tis_raw","tis_value","tis_unit","tis_ml","tis_status","ether_raw","ether_value","ether_unit","ether_ml","ether_status",
        "cocktail_total_ml","cocktail_ml_per_mmol","ether_ml_per_mmol","composition_pct_json","unit_review_required","source_file","source_page","source_locator","raw_note",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader()
        for row in rows: writer.writerow(row)


def test_v5_material_parser_preserves_unitless_values_for_review(tmp_path):
    from openpyxl import Workbook
    from suite_gui.v5_material_usage import extract_cleavage_usage_workbook

    wb=Workbook(); ws=wb.active; ws.title="usage"
    ws["A1"]="발주처& 펩타이드 명칭 :"; ws["B1"]="Client& SYNTH-USAGE-01"; ws["C1"]="scale: 0.4 mmol"
    ws["A2"]="서열 :"; ws["B2"]="WFRYKSRR"
    ws["A20"]="TFA(Trifluoroacetic acid)"; ws["B20"]="7.6 ml"
    ws["A21"]="H2O (cleavage)"; ws["B21"]=0.2
    ws["A22"]="TIS"; ws["B22"]="-"
    ws["A23"]="Ethyl Ether"; ws["B23"]="180 mL"
    source=tmp_path/"usage.xlsx"; wb.save(source)
    rows=extract_cleavage_usage_workbook(source)
    assert len(rows)==1
    row=rows[0]
    assert row["tfa_ml"]==7.6
    assert row["water_raw"]=="0.2" and row["water_ml"] is None
    assert row["water_status"]=="needs_unit_review"
    assert row["tis_ml"]==0.0
    assert row["ether_ml"]==180.0
    assert row["unit_review_required"]==1
    assert row["cocktail_total_ml"] is None


def test_v5_cleavage_amount_exact_and_bounded_interpolation(tmp_path):
    from suite_gui import experimental_data
    from suite_gui.decision_support_v5 import cleavage_amount_recommendation

    db=tmp_path/"exp.sqlite"; source=tmp_path/"usage.csv"
    base={"status":"parsed","product_raw":"X","product":"X","product_key":"x","sequence_raw":"ACDE","sequence":"ACDE","sequence_key":"|A|C|D|E|","operator":"","record_scope":"single","tfa_unit":"mL","tfa_status":"explicit_volume","water_unit":"mL","water_status":"explicit_volume","tis_unit":"","tis_status":"explicit_zero","ether_unit":"mL","ether_status":"explicit_volume","composition_pct_json":"{\"TFA\": 95.0, \"Water\": 5.0}","unit_review_required":0,"source_page":"usage","source_locator":"","raw_note":""}
    rows=[]
    for scale,total,ether,tag in [(1.0,20.0,100.0,"a"),(2.0,40.0,200.0,"b")]:
        rows.append({**base,"scale_mmol":scale,"manufacture_period":tag,"synthesis_key":f"x|acde|{tag}|{scale}","tfa_raw":str(total*.95),"tfa_value":total*.95,"tfa_ml":total*.95,"water_raw":str(total*.05),"water_value":total*.05,"water_ml":total*.05,"tis_raw":"-","tis_value":0,"tis_ml":0,"ether_raw":str(ether),"ether_value":ether,"ether_ml":ether,"cocktail_total_ml":total,"cocktail_ml_per_mmol":20.0,"ether_ml_per_mmol":100.0,"source_file":f"{tag}.xlsx"})
    _write_usage_csv(source,rows); experimental_data.import_cleavage_usage_csv(source,db)
    exact=cleavage_amount_recommendation(product="x",sequence="acde",scale_mmol=1,db_path=db)
    assert exact["recommendation_kind"]=="HISTORICAL EXACT SCALE"
    assert exact["recommended_amount"]["scaled_total_ml"]==20.0
    mid=cleavage_amount_recommendation(product="X",sequence="ACDE",scale_mmol=1.5,db_path=db)
    assert mid["recommendation_kind"]=="BOUNDED MODEL INTERPOLATION"
    assert mid["recommended_amount"]["scaled_total_ml"]==30.0
    outside=cleavage_amount_recommendation(product="X",sequence="ACDE",scale_mmol=3,db_path=db)
    assert outside["apply_allowed"] is False
    assert outside["recommended_amount"]["scaled_total_ml"] is None


def test_v5_difficulty_map_is_review_only():
    from suite_gui.decision_support_v5 import sequence_difficulty_map
    result=sequence_difficulty_map("Ac-D-G-V-V-M-C-NH2")
    assert result["positions"]
    assert result["apply_allowed"] is False
    assert result["score_kind"]=="deterministic review score"
    assert any("Aspartimide" in " ".join(row["reasons"]) for row in result["positions"])


def test_v5_outcomes_are_separate_from_conditions(tmp_path):
    from suite_gui import experimental_data
    db=tmp_path/"exp.sqlite"
    saved=experimental_data.add_outcome({"stage":"cleavage","product":"P","sequence":"ac-aaaa-nh2","result":"Success","success_flag":1,"yield_percent":71.2,"purity_percent":95.1,"observation":"ok"},db)
    assert saved["sequence_key"]==experimental_data.canonical_sequence_key("Ac-AAAA-NH2")
    rows=experimental_data.list_records("outcome",db)
    assert len(rows)==1 and rows[0]["yield_percent"]==71.2


def test_v5_model_rebuild_is_explicit_and_requires_verified_data(tmp_path, monkeypatch):
    from suite_gui import experimental_data, model_registry_v5
    monkeypatch.setattr(model_registry_v5,"_paths",lambda:(tmp_path/"registry.json",tmp_path/"model.joblib"))
    db=tmp_path/"exp.sqlite"
    result=model_registry_v5.rebuild_loading_model(db)
    assert result["built"] is False
    assert not (tmp_path/"model.joblib").exists()


def test_v5_stage_risk_advisor_separates_loading_coupling_cleavage(tmp_path):
    from suite_gui import experimental_data
    from suite_gui.decision_support_v5 import stage_risk_advisor
    db=tmp_path/'exp.sqlite'
    experimental_data.add_record('loading',{
        'resin_type':'Rink Amide AM','amino_acid_raw':'Fmoc-Trp(Boc)-OH','aa_eq':2,'base':'DIEA','base_eq':4,
        'loading_time_h':4,'loading_rate_mmol_g':0.71,
    },db,status='verified')
    experimental_data.add_outcome({'stage':'coupling','product':'P','sequence':'ACMW','result':'failed','success_flag':0},db,status='verified')
    result=stage_risk_advisor(sequence='ACMW',product='P',resin='Rink Amide AM',scale_mmol=1,db_path=db)
    by={row['stage']:row for row in result['stages']}
    assert set(by)=={'Loading','Coupling','Cleavage'}
    assert by['Loading']['evidence_count']==1
    assert by['Coupling']['level'] in {'HIGH','CRITICAL'}
    assert any('failures=1' in reason for reason in by['Coupling']['reasons'])
    assert by['Cleavage']['level'] in {'WARNING','MEDIUM','HIGH','CRITICAL'}
    assert result['apply_allowed'] is False
    assert 'not failure probability' in result['risk_kind']


def test_v5_loading_model_registry_keeps_versions_and_rolls_back(tmp_path, monkeypatch):
    from suite_gui import experimental_data, model_registry_v5
    monkeypatch.setattr(model_registry_v5,'_paths',lambda:(tmp_path/'registry.json',tmp_path/'loading_model_v5.joblib'))
    db=tmp_path/'exp.sqlite'
    aas=['A','V','R','G']
    for i in range(12):
        experimental_data.add_record('loading',{
            'resin_type':'2-CTC' if i%2 else 'Rink Amide AM','amino_acid_raw':aas[i%len(aas)],
            'aa_eq':1+(i%3),'base':'DIEA','base_eq':2+(i%2),'loading_time_h':2+(i%4),
            'loading_rate_mmol_g':0.35+0.04*(i%6),
        },db,status='verified')
    first=model_registry_v5.rebuild_loading_model(db)
    assert first['built'] is True
    second=model_registry_v5.rebuild_loading_model(db)
    assert second['built'] is True and second['model_id']!=first['model_id']
    history=model_registry_v5.loading_model_history()
    assert len(history)==2 and history[0]['active'] is True
    rolled=model_registry_v5.rollback_loading_model()
    assert rolled['activated'] is True and rolled['model_id']==first['model_id']
    assert model_registry_v5.loading_model_info()['active_model_id']==first['model_id']


def test_v5_worse_loading_model_is_kept_as_candidate_until_explicit_promotion(tmp_path, monkeypatch):
    from suite_gui import experimental_data, model_registry_v5
    import json
    monkeypatch.setattr(model_registry_v5,'_paths',lambda:(tmp_path/'registry.json',tmp_path/'loading_model_v5.joblib'))
    db=tmp_path/'exp.sqlite'
    aas=['A','V','R','G']
    for i in range(12):
        experimental_data.add_record('loading',{
            'resin_type':'2-CTC' if i%2 else 'Rink Amide AM','amino_acid_raw':aas[i%len(aas)],
            'aa_eq':1+(i%3),'base':'DIEA','base_eq':2+(i%2),'loading_time_h':2+(i%4),
            'loading_rate_mmol_g':0.35+0.04*(i%6),
        },db,status='verified')
    first=model_registry_v5.rebuild_loading_model(db)
    assert first['built'] is True and first['promoted'] is True
    registry=json.loads((tmp_path/'registry.json').read_text())
    registry['models']['loading']['versions'][0]['cross_validated_mae_mmol_g']=1e-9
    (tmp_path/'registry.json').write_text(json.dumps(registry),encoding='utf-8')
    second=model_registry_v5.rebuild_loading_model(db)
    assert second['built'] is True and second['promoted'] is False
    assert second['promotion_status']=='CANDIDATE'
    assert model_registry_v5.loading_model_info()['active_model_id']==first['model_id']
    promoted=model_registry_v5.promote_latest_loading_candidate()
    assert promoted['activated'] is True and promoted['model_id']==second['model_id']
    assert model_registry_v5.loading_model_info()['active_model_id']==second['model_id']
