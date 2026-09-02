from __future__ import annotations
import csv
from pathlib import Path

def _write_csv(path: Path, rows):
    fields=[]
    for row in rows:
        for k in row:
            if k not in fields: fields.append(k)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)

def test_v5_nhexane_is_distinct_workup_evidence(tmp_path):
    from openpyxl import Workbook
    from suite_gui import experimental_data
    from suite_gui.v5_material_usage import extract_cleavage_usage_workbook
    from suite_gui.decision_support_v5 import cleavage_amount_recommendation
    wb=Workbook(); ws=wb.active; ws.title='usage'
    ws['A1']='발주처& 펩타이드 명칭 :'; ws['B1']='Client& SYNTH-HEX-01'; ws['C1']='scale: 1 mmol'
    ws['A2']='서열 :'; ws['B2']='ACDE'
    ws['A20']='TFA(Trifluoroacetic acid)'; ws['B20']='19 mL'
    ws['A21']='H2O (cleavage)'; ws['B21']='1 mL'
    ws['A22']='TIS'; ws['B22']='-'
    ws['A23']='n-Hexane'; ws['B23']='100 mL'
    x=tmp_path/'hex.xlsx'; wb.save(x)
    rows=extract_cleavage_usage_workbook(x)
    assert rows[0]['hexane_ml']==100.0 and rows[0]['ether_ml'] is None
    db=tmp_path/'exp.sqlite'; experimental_data.import_cleavage_usage_workbook(x,db)
    result=cleavage_amount_recommendation(product='SYNTH-HEX-01',sequence='ACDE',scale_mmol=1,db_path=db)
    assert result['recommended_workup']['solvent']=='n-Hexane'
    assert result['recommended_workup']['scaled_workup_ml']==100.0

def test_v5_aggregate_rows_are_reference_only(tmp_path):
    from suite_gui import experimental_data
    from suite_gui.decision_support_v5 import cleavage_amount_recommendation
    db=tmp_path/'exp.sqlite'; src=tmp_path/'usage.csv'
    base=dict(status='parsed',product='X',product_raw='X',sequence='ACDE',sequence_raw='ACDE',scale_mmol=1,operator='',manufacture_period='',tfa_raw='19',tfa_value=19,tfa_unit='mL',tfa_ml=19,tfa_status='explicit_volume',water_raw='1',water_value=1,water_unit='mL',water_ml=1,water_status='explicit_volume',tis_raw='-',tis_value=0,tis_unit='',tis_ml=0,tis_status='explicit_zero',ether_raw='',ether_value='',ether_unit='',ether_ml='',ether_status='missing',hexane_raw='',hexane_value='',hexane_unit='',hexane_ml='',hexane_status='missing',cocktail_total_ml=20,cocktail_ml_per_mmol=20,composition_pct_json='{"TFA":95,"Water":5}',unit_review_required=0,source_page='usage',source_locator='',raw_note='')
    rows=[{**base,'synthesis_key':'single','record_scope':'single','source_file':'single.xlsx'},
          {**base,'synthesis_key':'aggregate','record_scope':'aggregate','cocktail_total_ml':200,'cocktail_ml_per_mmol':200,'source_file':'aggregate.xlsx'}]
    _write_csv(src,rows); experimental_data.import_cleavage_usage_csv(src,db)
    result=cleavage_amount_recommendation(product='X',sequence='ACDE',scale_mmol=1,db_path=db)
    assert result['recommended_amount']['scaled_total_ml']==20.0
    assert result['aggregate_reference_count']==1

def test_v5_visible_tabs_use_final_names():
    from pathlib import Path
    import suite_gui.modules.experimental_data_panel as panel
    text=Path(panel.__file__).read_text(encoding='utf-8')
    assert 'text="Risk & Evidence"' in text
    assert 'text="Recommended Conditions"' in text
    assert '1. Cleavage Condition' in text
    assert '2. Amount at Current Scale' in text
    assert '3. Precipitation / Workup' in text
